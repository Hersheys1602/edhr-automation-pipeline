from pathlib import Path
from datetime import datetime
import openpyxl


def clean_string(value):
    if value is None:
        return None
    value = str(value).strip()
    return value if value else None


def format_datetime(value):
    """
    Convert Excel datetime / date / time objects to ISO-like strings.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.isoformat(sep=" ")

    return str(value).strip()


def parse_mes_board_history(filepath: str) -> dict:
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"MES file not found: {filepath}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
    header_map = {header: idx + 1 for idx, header in enumerate(headers) if header is not None}

    def cell(row_num, header_name):
        col_num = header_map.get(header_name)
        if col_num is None:
            return None
        return ws.cell(row=row_num, column=col_num).value

    device_identification = {
        "serial_number": None,
        "assembly_id": None,
        "assembly_revision": None,
        "board_type": None,
        "board_revision": None,
        "image_id": None,   # MES usually won't populate this
        "customer": None,
        "manufacturing_start_time": None,
        "manufacturing_end_time": None,
    }

    process_execution_summary = []

    earliest_start = None
    latest_end = None

    for row_num in range(2, ws.max_row + 1):
        serial_number = clean_string(cell(row_num, "Serial Number"))
        assembly_raw = clean_string(cell(row_num, "Assembly"))
        customer = clean_string(cell(row_num, "Customer"))
        start_date = cell(row_num, "Start Date")
        end_date = cell(row_num, "End Date")
        military_time = cell(row_num, "Military Time")

        board_revision = (
            clean_string(cell(row_num, "Board Revision"))
            or clean_string(cell(row_num, "Board Rev"))
            or clean_string(cell(row_num, "Revision"))
        )

        if device_identification["serial_number"] is None and serial_number:
            device_identification["serial_number"] = serial_number

        if device_identification["assembly_id"] is None and assembly_raw:
            device_identification["assembly_id"] = assembly_raw

            # Example expected pattern:
            # "10081656-02 / 02 /  / Polaris FF2 PCBA"
            parts = [p.strip() for p in assembly_raw.split("/")]

            if parts and parts[0]:
                device_identification["assembly_id"] = parts[0]

            if len(parts) >= 2 and parts[1]:
                device_identification["assembly_revision"] = parts[1]

            if len(parts) >= 4 and parts[3]:
                device_identification["board_type"] = parts[3]

        if device_identification["customer"] is None and customer:
            device_identification["customer"] = customer

        if device_identification["board_revision"] is None and board_revision:
            device_identification["board_revision"] = board_revision

        if isinstance(start_date, datetime):
            if earliest_start is None or start_date < earliest_start:
                earliest_start = start_date

        if isinstance(end_date, datetime):
            if latest_end is None or end_date > latest_end:
                latest_end = end_date

        row_data = {
            "route_name": clean_string(cell(row_num, "Factory / Route / MA")),
            "process_step": clean_string(cell(row_num, "Route Step")),
            "equipment_id": clean_string(cell(row_num, "Equipment")),
            "system_version": clean_string(cell(row_num, "CRD")),
            "operator_id": clean_string(cell(row_num, "Operator Name")),
            "military_time": format_datetime(military_time),
            "status": clean_string(cell(row_num, "Status")),
        }

        if any(value not in [None, "", " "] for value in row_data.values()):
            process_execution_summary.append(row_data)

    device_identification["manufacturing_start_time"] = format_datetime(earliest_start)
    device_identification["manufacturing_end_time"] = format_datetime(latest_end)

    return {
        "device_identification": device_identification,
        "process_execution_summary": process_execution_summary,
    }