import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT


def prettify_key(key: str) -> str:
    return key.replace('_', ' ').title()


def add_title(doc: Document, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def add_field_block(doc: Document, data: dict, skip_empty: bool = False):
    if not data:
        doc.add_paragraph("No data available.")
        return

    for key, value in data.items():
        if skip_empty and (value is None or value == "" or value == [] or value == {}):
            continue

        p = doc.add_paragraph()
        label_run = p.add_run(f"{prettify_key(key)}: ")
        label_run.bold = True
        p.add_run("" if value is None else str(value))


def add_two_column_table(doc: Document, title: str, data: dict, skip_empty: bool = False):
    add_heading(doc, title, level=2)

    if not data or all(v in [None, "", [], {}] for v in data.values()):
        doc.add_paragraph("No data available.")
        doc.add_paragraph()
        return

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Field"
    hdr_cells[1].text = "Value"

    for key, value in data.items():
        if skip_empty and (value is None or value == "" or value == [] or value == {}):
            continue

        row_cells = table.add_row().cells
        row_cells[0].text = prettify_key(key)
        row_cells[1].text = "" if value is None else str(value)

    doc.add_paragraph()


def add_list_of_dicts_table(doc: Document, title: str, rows: list[dict], skip_empty_rows: bool = True):
    add_heading(doc, title, level=2)

    if not rows:
        doc.add_paragraph("No data available.")
        doc.add_paragraph()
        return

    if skip_empty_rows:
        filtered_rows = []
        for row in rows:
            if any(v not in [None, "", [], {}] for v in row.values()):
                filtered_rows.append(row)
        rows = filtered_rows

    if not rows:
        doc.add_paragraph("No data available.")
        doc.add_paragraph()
        return

    headers = list(rows[0].keys())

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = prettify_key(header)

    for row in rows:
        row_cells = table.add_row().cells
        for i, header in enumerate(headers):
            value = row.get(header)
            row_cells[i].text = "" if value is None else str(value)

    doc.add_paragraph()


def render_edhr_docx(json_path: str, output_path: str):
    json_file = Path(json_path)
    if not json_file.exists():
        raise FileNotFoundError(f"JSON file not found: {json_path}")

    with open(json_file, "r", encoding="utf-8") as f:
        edhr = json.load(f)

    # guardrail: this renderer expects the new sectioned AWS-assembled schema
    required_top = {
        "document_control",
        "device_identification",
        "process_execution_summary",
        "inspection_and_defects",
        "rework_documentation",
        "deviation_documentation",
        "component_traceability",
        "electrical_test_certification",
    }
    missing = [k for k in required_top if k not in edhr]
    if missing:
        raise ValueError(
            "This renderer expects the NEW section-based AWS eDHR JSON. "
            f"Missing top-level keys: {missing}"
        )

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    add_title(doc, "Electronic Device History Record (eDHR)")

    add_heading(doc, "1. Document Control", level=1)
    add_field_block(doc, edhr.get("document_control", {}), skip_empty=False)
    doc.add_paragraph()

    add_heading(doc, "2. Device Identification", level=1)
    add_field_block(doc, edhr.get("device_identification", {}), skip_empty=False)
    doc.add_paragraph()

    add_heading(doc, "3. Process Execution Summary", level=1)

    # get rows
    rows = edhr.get("process_execution_summary", [])

# enforce column order if rows exist
    if rows:
        ordered_keys = [
            "route_name",
            "process_step",
            "equipment_id",
            "system_version",
            "operator_id",
            "military_time",
            "status",
        ]

    reordered_rows = []
    for row in rows:
        reordered = {k: row.get(k) for k in ordered_keys}
        reordered_rows.append(reordered)

    rows = reordered_rows

# now render
    add_list_of_dicts_table(
    doc,
    title="Process Execution Summary",
    rows=rows,
)

    add_heading(doc, "4. Inspection & Defects", level=1)
    inspection = edhr.get("inspection_and_defects", {})

    add_list_of_dicts_table(
        doc,
        title="Defect Summary",
        rows=inspection.get("defect_summary", []),
    )

    add_two_column_table(
        doc,
        title="AOI Inspection Summary",
        data=inspection.get("aoi_inspection_summary", {}),
        skip_empty=True,
    )

    add_two_column_table(
        doc,
        title="AXI Inspection Summary",
        data=inspection.get("axi_inspection_summary", {}),
        skip_empty=True,
    )

    add_two_column_table(
        doc,
        title="Flying Probe Inspection Summary",
        data=inspection.get("flying_probe_inspection_summary", {}),
        skip_empty=True,
    )

    add_heading(doc, "5. Rework Documentation", level=1)
    add_list_of_dicts_table(
        doc,
        title="Rework Documentation",
        rows=edhr.get("rework_documentation", []),
    )

    add_heading(doc, "6. Deviation Documentation", level=1)
    add_list_of_dicts_table(
        doc,
        title="Deviation Documentation",
        rows=edhr.get("deviation_documentation", []),
    )

    add_heading(doc, "7. Component Traceability", level=1)
    add_list_of_dicts_table(
        doc,
        title="Component Traceability",
        rows=edhr.get("component_traceability", []),
    )

    add_heading(doc, "8. Electrical Test Certification", level=1)
    add_two_column_table(
        doc,
        title="Electrical Test Certification",
        data=edhr.get("electrical_test_certification", {}),
        skip_empty=True,
    )

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    pass
  
